# Copy of ERPNext "General Ledger" report.
# Only difference: from OPENING_DATE onwards, the opening balance is taken from
# OPENING_BALANCES below instead of the GL Entries posted before that date.

import re
from datetime import timedelta
from io import BytesIO

import frappe
from frappe import _, _dict
from frappe.utils import cstr, flt, getdate

from erpnext.accounts.doctype.accounting_dimension.accounting_dimension import get_accounting_dimensions
from erpnext.accounts.report.general_ledger.general_ledger import (
	get_accountwise_gle,
	get_columns,
	get_data_with_opening_closing,
	get_gl_entries,
	get_result_as_list,
	initialize_gle_map,
	set_bill_no,
	set_account_currency,
	update_translations,
	validate_filters,
	validate_party,
)

OPENING_DATE = "2025-01-01"

# Opening balances as on OPENING_DATE: {company: {account_number: (debit, credit)}}
# Accounts not listed here start from zero on OPENING_DATE.
OPENING_BALANCES = {
	"Elite Resources Center": {
		"11010201": (779741, 0),  # Alinma Bank
		"11010202": (1747997, 0),  # AlJazira Bank 901
		"11010204": (5161, 0),  # Alinma Bank 5003
		"11010205": (6, 0),  # AlJazira Bank 902
		"11040101": (37125, 0),  # Prepaid Rent
		"11030101": (7977324, 0),  # Accounts receivable
		"11030401": (77209, 0),  # Employee Advances
		"2230101": (0, 50116),  # Accrued Expenses - Salaries
		"2230102": (0, 518263),  # Accrued Expenses - GOSI
		"2230103": (0, 18000),  # Accrued Expenses - Audit Fees
		"22304": (0, 4924387),  # VAT payable
		"2230109-01": (0, 1110150),  # ESB
		"21203": (0, 364236),  # Zakat Provision
		"311": (0, 250000),  # Capital
		"321": (10088104, 0),  # Owner's Current Account
		"331": (0, 13477515),  # Retained Earnings
	}
}


def execute(filters=None):
	if not filters:
		return [], []

	filters, account_details = prepare_filters(filters)

	columns = get_columns(filters)

	update_translations()

	res = get_result(filters, account_details)

	return columns, res


def prepare_filters(filters):
	account_details = {}

	if filters and filters.get("print_in_account_currency") and not filters.get("account"):
		frappe.throw(_("Select an account to print in account currency"))

	for acc in frappe.db.sql("""select name, is_group from tabAccount""", as_dict=1):
		account_details.setdefault(acc.name, acc)

	if filters.get("party"):
		filters.party = frappe.parse_json(filters.get("party"))

	validate_filters(filters, account_details)

	validate_party(filters)

	filters = set_account_currency(filters)

	return filters, account_details


def get_result(filters, account_details):
	accounting_dimensions, gl_entries = get_entries(filters)

	data = get_data_with_opening_closing(filters, account_details, accounting_dimensions, gl_entries)

	result = get_result_as_list(data, filters)

	return result


def get_entries(filters):
	accounting_dimensions = []
	if filters.get("include_dimensions"):
		accounting_dimensions = get_accounting_dimensions()

	gl_entries = get_gl_entries(filters, accounting_dimensions)

	if use_custom_opening(filters):
		gl_entries = apply_custom_opening(filters, gl_entries)

	return accounting_dimensions, gl_entries


def use_custom_opening(filters):
	return (
		filters.get("company") in OPENING_BALANCES
		and getdate(filters.from_date) >= getdate(OPENING_DATE)
		and not filters.get("party")
	)


def apply_custom_opening(filters, gl_entries):
	opening_date = getdate(OPENING_DATE)

	# drop everything posted before OPENING_DATE (the wrong old opening)
	gl_entries = [gle for gle in gl_entries if getdate(gle.posting_date) >= opening_date]

	# filters.account is already expanded with child accounts by get_gl_entries
	selected_accounts = set(filters.account) if filters.get("account") else None

	opening_entries = []
	posting_date = opening_date - timedelta(days=1)
	for account_number, (debit, credit) in OPENING_BALANCES[filters.company].items():
		account = frappe.db.get_value(
			"Account", {"company": filters.company, "account_number": account_number}, "name"
		)
		if not account or (selected_accounts is not None and account not in selected_accounts):
			continue
		if not flt(debit) and not flt(credit):
			continue

		opening_entries.append(
			_dict(
				gl_entry=None,
				posting_date=posting_date,
				account=account,
				party_type=None,
				party=None,
				voucher_type=None,
				voucher_no=None,
				cost_center=None,
				project=None,
				against_voucher_type=None,
				against_voucher=None,
				account_currency=None,
				against=None,
				is_opening="No",
				debit=flt(debit),
				credit=flt(credit),
				debit_in_account_currency=flt(debit),
				credit_in_account_currency=flt(credit),
			)
		)

	return opening_entries + gl_entries


@frappe.whitelist()
def export_sheet_per_account(filters):
	"""Download the report as one .xlsx: a Summary sheet, then one sheet per account
	with the same columns and Opening / entries / Total / Closing rows as the report."""
	from openpyxl import Workbook
	from openpyxl.styles import Font

	if not frappe.get_doc("Report", "GL").is_permitted():
		frappe.throw(_("Not permitted"), frappe.PermissionError)

	filters = _dict(frappe.parse_json(filters) or {})
	# validate as consolidated: ERPNext rejects an Account filter with Group by Account
	filters.group_by = "Group by Voucher (Consolidated)"
	filters, _account_details = prepare_filters(filters)
	filters.group_by = "Group by Account"
	update_translations()
	columns = [c for c in get_columns(filters) if not c.get("hidden")]

	# Same as the report's Group by Account, but an account with an opening and no
	# entries in the period is kept too (the report skips it), so the sheets add up.
	accounting_dimensions, gl_entries = get_entries(filters)
	set_bill_no(gl_entries)
	gle_map = initialize_gle_map(gl_entries, filters)
	totals, _entries = get_accountwise_gle(filters, accounting_dimensions, gl_entries, gle_map)

	blocks = []
	for account, acc_dict in gle_map.items():
		t = acc_dict.totals
		if not acc_dict.entries and not any(flt(t.opening.get(k), 2) for k in ("debit", "credit")):
			continue
		block = get_result_as_list([t.opening, *acc_dict.entries, t.total, t.closing], filters)
		blocks.append((account, block))
	grand_total, grand_closing = totals.total, totals.closing

	bold = Font(bold=True)
	wb = Workbook()
	summary = wb.active
	summary.title = "Summary"
	summary.append([_("Account"), _("Opening Debit"), _("Opening Credit"), _("Debit"), _("Credit"),
		_("Closing Debit"), _("Closing Credit"), _("Balance")])
	for cell in summary[1]:
		cell.font = bold

	used_names = {"summary"}
	for account, block in blocks:
		opening, total, closing = block[0], block[-2], block[-1]
		summary.append([account, flt(opening.get("debit")), flt(opening.get("credit")),
			flt(total.get("debit")), flt(total.get("credit")),
			flt(closing.get("debit")), flt(closing.get("credit")),
			flt(closing.get("debit")) - flt(closing.get("credit"))])

		ws = wb.create_sheet(_sheet_name(account, used_names))
		ws.append([_("Account"), account])
		ws.append([_("Period"), f"{filters.from_date} to {filters.to_date}"])
		ws.append([])
		ws.append([c.get("label") for c in columns])
		for cell in ws[4]:
			cell.font = bold
		for row in block:
			ws.append([_cell_value(row.get(c.get("fieldname")), c) for c in columns])
			if not row.get("posting_date"):
				for cell in ws[ws.max_row]:
					cell.font = bold
		_format_sheet(ws, columns, header_row=4)

	summary.append([])
	summary.append([cstr(grand_total.get("account")).strip("'"), None, None,
		flt(grand_total.get("debit")), flt(grand_total.get("credit"))])
	summary.append([cstr(grand_closing.get("account")).strip("'"), None, None, None, None,
		flt(grand_closing.get("debit")), flt(grand_closing.get("credit")),
		flt(grand_closing.get("debit")) - flt(grand_closing.get("credit"))])
	for row in summary.iter_rows(min_row=summary.max_row - 1):
		for cell in row:
			cell.font = bold
	summary.column_dimensions["A"].width = 45
	for col in "BCDEFGH":
		summary.column_dimensions[col].width = 16
	for row in summary.iter_rows(min_row=2, min_col=2):
		for cell in row:
			cell.number_format = "#,##0.00"
	summary.freeze_panes = "A2"

	out = BytesIO()
	wb.save(out)
	frappe.response["filename"] = f"GL {filters.from_date} to {filters.to_date}.xlsx"
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


def _sheet_name(account, used_names):
	# Excel: max 31 chars, unique (case-insensitive), none of []:*?/\
	base = re.sub(r"[\[\]:*?/\\]", " ", cstr(account)).strip()[:31] or "Account"
	name, i = base, 2
	while name.lower() in used_names:
		suffix = f" ({i})"
		name, i = base[: 31 - len(suffix)] + suffix, i + 1
	used_names.add(name.lower())
	return name


def _cell_value(value, column):
	if value is None:
		return None
	if column.get("fieldtype") in ("Float", "Currency"):
		return flt(value)
	if column.get("fieldtype") == "Date":
		return getdate(value)
	if column.get("fieldname") == "account":
		return cstr(value).strip("'")
	return cstr(value)


def _format_sheet(ws, columns, header_row):
	from openpyxl.utils import get_column_letter

	for i, column in enumerate(columns, 1):
		letter = get_column_letter(i)
		ws.column_dimensions[letter].width = max(12, min(45, (column.get("width") or 100) / 7))
		fmt = {"Float": "#,##0.00", "Currency": "#,##0.00", "Date": "yyyy-mm-dd"}.get(column.get("fieldtype"))
		if fmt:
			for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=i, max_col=i):
				cell.number_format = fmt
	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
